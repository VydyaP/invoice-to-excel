import csv
import datetime as dt
import io

from openpyxl import load_workbook

from conftest import KeyProvider
from invoice_to_excel.export import INVOICE_COLUMNS, LINE_COLUMNS, invoice_rows, line_rows, to_csv, to_excel
from invoice_to_excel.pipeline import estimated_cost_usd, process, recheck
from invoice_to_excel.preview import preview_images
from invoice_to_excel.schema import invoice_from_dict

TODAY = dt.date(2026, 10, 31)


def run(sample_files, **kw):
    return process(sample_files, KeyProvider(**kw), TODAY)


def test_statuses_for_the_sample_batch(sample_files):
    results = run(sample_files)
    status = {r.file_name: r.status for r in results}
    assert status["inv01_sahyadri.pdf"] == "Looks right"
    assert status["inv08_konkan_2026.pdf"] == "Check warnings"
    assert status["inv09_sahyadri_badtotal.pdf"] == "Needs review"
    assert sum(s == "Looks right" for s in status.values()) == 7


def test_one_failed_file_does_not_stop_the_batch(sample_files):
    results = run(sample_files, fail={"inv03_coromandel.pdf"})
    bad = next(r for r in results if r.file_name == "inv03_coromandel.pdf")
    assert bad.status == "Could not read" and bad.invoice is None and "temporary" in bad.error
    assert sum(r.invoice is not None for r in results) == 11


def test_fixing_a_field_clears_its_flag(sample_files):
    results = run(sample_files)
    r = next(r for r in results if r.file_name == "inv09_sahyadri_badtotal.pdf")
    data = r.invoice.to_dict()
    data["grand_total"] = 7564.0
    r.invoice = invoice_from_dict(data)
    recheck(results, TODAY)
    assert r.status == "Looks right"


def test_fixing_the_original_moves_duplicate_flag(sample_files):
    results = run(sample_files)
    dup = next(r for r in results if r.file_name == "inv12_sahyadri_duplicate.jpg")
    data = dup.invoice.to_dict()
    data["invoice_number"] = "SAF/26-27/0413"
    dup.invoice = invoice_from_dict(data)
    recheck(results, TODAY)
    assert not any(f.code == "DUPLICATE" for f in dup.flags)


def test_cost_estimate(sample_files):
    results = run(sample_files)
    assert abs(estimated_cost_usd(results) - 12 * (3000 * 0.10 + 600 * 0.50) / 1e6) < 1e-9
    assert estimated_cost_usd(run(sample_files, model="unknown-model")) is None


def test_excel_workbook(sample_files):
    results = run(sample_files, fail={"inv03_coromandel.pdf"})
    wb = load_workbook(io.BytesIO(to_excel(results, "claude-haiku-5-5")))
    assert wb.sheetnames == ["Invoices", "Line items", "Checks", "About"]
    inv = wb["Invoices"]
    header = [c.value for c in inv[1]]
    assert header == [c[0] for c in INVOICE_COLUMNS] and inv.max_row == 13
    row = {h: c for h, c in zip(header, inv[2], strict=True)}
    assert row["File"].value == "inv01_sahyadri.pdf"
    assert row["Invoice date"].value == dt.datetime(2026, 10, 3)
    assert "##,##0.00" in row["Grand total"].number_format and "₹" in row["Grand total"].number_format
    assert isinstance(row["Grand total"].value, (int, float))  # a number, not text
    checks = [[c.value for c in r] for r in wb["Checks"].iter_rows(min_row=2)]
    assert ["inv03_coromandel.pdf", "error", "NOT_READ"] == checks[0][:3] or any(c[2] == "NOT_READ" for c in checks)
    assert {c[2] for c in checks} >= {"TOTAL_MISMATCH", "GSTIN_CHECKSUM", "TAX_TYPE", "DUPLICATE", "RATE_NOT_VALID_ON_DATE"}
    lines = sum(len(r.invoice.line_items) for r in results if r.invoice)
    assert wb["Line items"].max_row == lines + 1
    about = {r[0].value: r[1].value for r in wb["About"].iter_rows()}
    assert about["Read by"] == "claude-haiku-5-5" and about["Needing review"] == 5


def test_csv_has_bom_and_rupee_safe(sample_files):
    results = run(sample_files)
    raw = to_csv(invoice_rows(results), INVOICE_COLUMNS)
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0][0] == "File" and len(rows) == 13
    lines = list(csv.reader(io.StringIO(to_csv(line_rows(results), LINE_COLUMNS).decode("utf-8-sig"))))
    assert lines[1][3]  # description present


def test_preview_for_pdf_photo_and_junk(sample_files):
    files = dict(sample_files)
    assert len(preview_images("inv01_sahyadri.pdf", files["inv01_sahyadri.pdf"])) == 1
    photo = preview_images("inv05_malnad_photo.jpg", files["inv05_malnad_photo.jpg"])
    assert len(photo) == 1 and photo[0].width <= 1100
    assert preview_images("x.pdf", b"broken") == []
