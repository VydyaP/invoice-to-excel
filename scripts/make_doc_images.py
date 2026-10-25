"""Build README images from the sample invoices.

    python scripts/make_doc_images.py

* docs/images/samples.png : a sample PDF invoice beside a phone photo of one
* docs/images/excel.png   : the Invoices sheet of the exported workbook
                            (needs LibreOffice `soffice` and `pdftoppm`; skipped if missing)

The Excel image uses the answer keys (the correct values), so it shows the export
format and the checks, not AI accuracy. docs/images/review.png is a screenshot of the
running app.
"""

from __future__ import annotations

import datetime as dt
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from invoice_to_excel.export import to_excel  # noqa: E402
from invoice_to_excel.extract import Extraction  # noqa: E402
from invoice_to_excel.pipeline import process  # noqa: E402
from invoice_to_excel.preview import preview_images  # noqa: E402

SAMPLES = ROOT / "samples" / "invoices"
ANSWERS = ROOT / "samples" / "answers"
IMAGES = ROOT / "docs" / "images"


def font(size: int):
    import matplotlib

    path = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans-Bold.ttf"
    return ImageFont.truetype(str(path), size)


def samples_image() -> None:
    pdf = preview_images("x.pdf", (SAMPLES / "inv01_sahyadri.pdf").read_bytes(), width=900)[0]
    photo = preview_images("x.jpg", (SAMPLES / "inv05_malnad_photo.jpg").read_bytes(), width=900)[0]
    h = 1150
    tiles = [im.resize((round(im.width * h / im.height), h)) for im in (pdf.crop((0, 0, pdf.width, min(pdf.height, round(pdf.width * 1.28)))), photo)]
    gap, top = 40, 70
    canvas = Image.new("RGB", (sum(t.width for t in tiles) + gap * 3, h + top + gap), "#F4F6F8")
    draw = ImageDraw.Draw(canvas)
    x = gap
    for tile, label in zip(tiles, ["PDF from billing software", "Phone photo of a printed bill"], strict=True):
        draw.text((x, 22), label, fill="#1F3B5A", font=font(30))
        canvas.paste(tile, (x, top))
        x += tile.width + gap
    canvas.save(IMAGES / "samples.png", optimize=True)


class AnswerKeys:
    name = "answer-key"

    def extract(self, file_name: str, data: bytes) -> Extraction:
        return Extraction(json.loads((ANSWERS / f"{Path(file_name).stem}.json").read_text("utf-8"))["invoice"], "answer key")


def excel_image() -> None:
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        print("soffice/pdftoppm not found: skipping excel.png")
        return
    from openpyxl import load_workbook

    files = [(p.name, b"") for p in sorted(SAMPLES.iterdir())]
    results = process(files, AnswerKeys(), dt.date(2026, 10, 31))
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        book = tmp / "invoices.xlsx"
        book.write_bytes(to_excel(results, "answer key"))
        wb = load_workbook(book)
        for name in wb.sheetnames[1:]:
            del wb[name]
        ws = wb["Invoices"]
        for col in ("D", "E", "G", "H", "M"):  # hide a few columns so the picture stays readable
            ws.column_dimensions[col].hidden = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.save(book)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(book)],
                       check=True, capture_output=True, timeout=180)
        subprocess.run(["pdftoppm", "-png", "-r", "150", "-singlefile", str(tmp / "invoices.pdf"), str(tmp / "page")],
                       check=True, timeout=60)
        img = Image.open(tmp / "page.png").convert("RGB")
        bg = Image.new("RGB", img.size, "white")
        from PIL import ImageChops

        box = ImageChops.difference(img, bg).getbbox()
        img.crop((box[0] - 20, box[1] - 20, box[2] + 20, box[3] + 20)).save(IMAGES / "excel.png", optimize=True)


if __name__ == "__main__":
    IMAGES.mkdir(parents=True, exist_ok=True)
    samples_image()
    excel_image()
    print("Images written to", IMAGES.relative_to(ROOT))
