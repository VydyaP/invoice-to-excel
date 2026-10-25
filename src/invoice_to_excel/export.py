"""Excel and CSV output. One row per invoice, one row per line item, and every check."""

from __future__ import annotations

import csv
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import Result

INR = '[>=10000000]"₹"##\\,##\\,##\\,##0.00;[>=100000]"₹"##\\,##\\,##0.00;"₹"##,##0.00'
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3B5A")
STATUS_FILL = {"Looks right": "E3F4E8", "Check warnings": "FFF4CE", "Needs review": "FDE2E1",
               "Could not read": "FDE2E1"}

INVOICE_COLUMNS = [
    ("File", "file", None, 28), ("Status", "status", None, 15), ("Supplier", "supplier_name", None, 32),
    ("Supplier GSTIN", "supplier_gstin", None, 18), ("Invoice no.", "invoice_number", None, 18),
    ("Invoice date", "invoice_date", "dd-mm-yyyy", 13), ("Buyer GSTIN", "buyer_gstin", None, 18),
    ("Place of supply", "place_of_supply_state_code", None, 9), ("Taxable value", "taxable_total", INR, 15),
    ("CGST", "cgst_amount", INR, 12), ("SGST", "sgst_amount", INR, 12), ("IGST", "igst_amount", INR, 12),
    ("Round off", "round_off", "0.00", 10), ("Grand total", "grand_total", INR, 15), ("Issues", "issues", "0", 8),
]
LINE_COLUMNS = [
    ("File", "file", None, 28), ("Supplier GSTIN", "supplier_gstin", None, 18),
    ("Invoice no.", "invoice_number", None, 18),
    ("Description", "description", None, 34), ("HSN/SAC", "hsn_code", None, 10), ("Qty", "quantity", "#,##0.###", 9),
    ("Unit", "unit", None, 7), ("Rate", "rate", INR, 12), ("Taxable value", "taxable_value", INR, 14),
    ("GST %", "gst_rate", "0.##", 7),
]
CHECK_COLUMNS = [("File", "file", None, 28), ("Severity", "severity", None, 10), ("Check", "code", None, 24),
                 ("Field", "field", None, 16), ("What to check", "message", None, 90)]


def invoice_rows(results: list[Result]) -> list[dict]:
    rows = []
    for r in results:
        inv = r.invoice
        base = {"file": r.file_name, "status": r.status, "issues": len(r.flags) + (1 if r.error else 0)}
        if inv:
            base.update({k: getattr(inv, k) for k in (
                "supplier_name", "supplier_gstin", "invoice_number", "invoice_date", "buyer_gstin",
                "place_of_supply_state_code", "taxable_total", "cgst_amount", "sgst_amount", "igst_amount",
                "round_off", "grand_total")})
        rows.append(base)
    return rows


def line_rows(results: list[Result]) -> list[dict]:
    rows = []
    for r in results:
        if not r.invoice:
            continue
        for ln in r.invoice.line_items:
            rows.append({"file": r.file_name, "supplier_gstin": r.invoice.supplier_gstin,
                         "invoice_number": r.invoice.invoice_number, **vars(ln)})
    return rows


def check_rows(results: list[Result]) -> list[dict]:
    rows = []
    for r in results:
        if r.error:
            rows.append({"file": r.file_name, "severity": "error", "code": "NOT_READ", "field": "", "message": r.error})
        for f in r.flags:
            rows.append({"file": r.file_name, "severity": f.severity, "code": f.code, "field": f.field,
                         "message": f.message})
    return rows


def _excel_value(key: str, value):
    if key == "invoice_date" and isinstance(value, str):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _sheet(wb: Workbook, title: str, columns, rows, first: bool = False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for j, (label, _, _, width) in enumerate(columns, 1):
        c = ws.cell(1, j, label)
        c.font, c.fill = HEADER, HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(rows, 2):
        for j, (_, key, fmt, _) in enumerate(columns, 1):
            value = _excel_value(key, row.get(key))
            c = ws.cell(i, j, value)
            if fmt and value is not None:
                c.number_format = fmt
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    return ws


def to_excel(results: list[Result], model: str | None = None) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Invoices", INVOICE_COLUMNS, invoice_rows(results), first=True)
    status_col = [c[1] for c in INVOICE_COLUMNS].index("status") + 1
    for i in range(2, ws.max_row + 1):
        fill = STATUS_FILL.get(ws.cell(i, status_col).value)
        if fill:
            ws.cell(i, status_col).fill = PatternFill("solid", fgColor=fill)
    _sheet(wb, "Line items", LINE_COLUMNS, line_rows(results))
    checks = _sheet(wb, "Checks", CHECK_COLUMNS, check_rows(results))
    for i in range(2, checks.max_row + 1):
        checks.cell(i, 5).alignment = Alignment(wrap_text=True, vertical="top")
    about = wb.create_sheet("About")
    lines = [
        ("Created", dt.datetime.now().strftime("%d %b %Y %H:%M")),
        ("Read by", model or "AI"),
        ("Invoices", len(results)),
        ("Needing review", sum(1 for r in results if r.status in ("Needs review", "Could not read"))),
        ("Note", "Fields were read by AI and checked by rules. Review the Checks sheet before posting to accounts."),
    ]
    for i, (k, v) in enumerate(lines, 1):
        about.cell(i, 1, k).font = Font(bold=True)
        about.cell(i, 2, v)
    about.column_dimensions["A"].width = 16
    about.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(rows: list[dict], columns) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in columns])
    for row in rows:
        w.writerow(["" if row.get(c[1]) is None else row.get(c[1]) for c in columns])
    # UTF-8 with BOM so Excel on Windows shows ₹ and names correctly
    return buf.getvalue().encode("utf-8-sig")
